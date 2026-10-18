"""Build a .pptx deck from a JSON spec with python-pptx.

Usage:
    python build_deck.py spec.json            # writes spec["output"]
    python build_deck.py spec.json -o out.pptx

The script never overwrites an existing file; pick a new name (e.g. add a date
suffix) instead. See SKILL.md for the spec format.
"""

import argparse
import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
CONTENT_TOP, CONTENT_BOTTOM = 1.65, 6.85

THEMES = {
    "blue": {"primary": "1F4E79", "accent": "2E75B6", "text": "262626", "muted": "7F7F7F", "light": "EAF1F8"},
    "dark": {"primary": "1E1E2E", "accent": "F2A541", "text": "262626", "muted": "7F7F7F", "light": "F3F0EA"},
    "green": {"primary": "1D5C4B", "accent": "3FA37C", "text": "262626", "muted": "7F7F7F", "light": "E8F4EE"},
}
DEFAULT_FONT = "맑은 고딕"


# ---------------------------------------------------------------- helpers

def rgb(hex_str):
    return RGBColor.from_string(hex_str.lstrip("#").upper())


def set_typeface(font, name):
    """Set Latin, East Asian and complex-script typefaces.

    font.name only sets <a:latin>, so Korean text would fall back to the theme
    font. Writing <a:ea> as well makes Hangul render in the requested font.
    """
    rPr = font._rPr
    for tag in ("a:latin", "a:ea", "a:cs"):
        old = rPr.find(qn(tag))
        if old is not None:
            rPr.remove(old)
    anchor = None
    for tag in ("a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"):
        anchor = rPr.find(qn(tag))
        if anchor is not None:
            break
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.makeelement(qn(tag), {"typeface": name})
        if anchor is not None:
            anchor.addprevious(el)
        else:
            rPr.append(el)


def style_font(font, ctx, size, color=None, bold=False, italic=False):
    font.size = Pt(size)
    font.bold = bold
    font.italic = italic
    font.color.rgb = rgb(color or ctx["theme"]["text"])
    set_typeface(font, ctx["font"])


def add_text(slide, ctx, text, left, top, width, height, size, color=None, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        style_font(run.font, ctx, size, color, bold)
    return box


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def set_bullet(paragraph, level, char="•"):
    pPr = paragraph._p.get_or_add_pPr()
    indent = 0.3
    pPr.set("marL", str(Emu(Inches(indent * (level + 1)))))
    pPr.set("indent", str(-Emu(Inches(indent))))
    for tag in ("a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip"):
        old = pPr.find(qn(tag))
        if old is not None:
            pPr.remove(old)
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}))


def flatten_bullets(items, level=0):
    """["a", ["a1", "a2"], "b"] -> [(0, "a"), (1, "a1"), (1, "a2"), (0, "b")]"""
    out = []
    for item in items:
        if isinstance(item, list):
            out.extend(flatten_bullets(item, level + 1))
        else:
            out.append((level, str(item)))
    return out


def fill_bullets(box, ctx, items, size):
    tf = box.text_frame
    tf.word_wrap = True
    colors = [ctx["theme"]["text"], ctx["theme"]["muted"]]
    for i, (level, text) in enumerate(flatten_bullets(items)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_bullet(p, level, "•" if level == 0 else "–")
        p.space_before = Pt(10 if level == 0 else 4)
        run = p.add_run()
        run.text = text
        style_font(run.font, ctx, size - 2 * min(level, 2), colors[min(level, 1)])


def bullet_size(items):
    n = len(flatten_bullets(items))
    return max(14, 22 - max(0, n - 5) * 1.5)


# ---------------------------------------------------------------- slide parts

def header(slide, ctx, title):
    add_text(slide, ctx, title, MARGIN, 0.45, SLIDE_W - 2 * MARGIN, 0.8, 28,
             ctx["theme"]["primary"], bold=True, anchor=MSO_ANCHOR.BOTTOM)
    add_rect(slide, MARGIN, 1.3, 1.2, 0.06, ctx["theme"]["accent"])


def footer(slide, ctx, number):
    if ctx.get("footer"):
        add_text(slide, ctx, ctx["footer"], MARGIN, 7.0, 8, 0.3, 10, ctx["theme"]["muted"])
    add_text(slide, ctx, str(number), SLIDE_W - MARGIN - 1, 7.0, 1, 0.3, 10,
             ctx["theme"]["muted"], align=PP_ALIGN.RIGHT)


def cover(slide, ctx, spec, size):
    t = ctx["theme"]
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, t["primary"])
    add_rect(slide, MARGIN + 0.05, 3.55, 1.5, 0.08, t["accent"])
    add_text(slide, ctx, spec.get("title", ""), MARGIN, 1.6, SLIDE_W - 2 * MARGIN, 1.85, size,
             "FFFFFF", bold=True, anchor=MSO_ANCHOR.BOTTOM)
    if spec.get("subtitle"):
        add_text(slide, ctx, spec["subtitle"], MARGIN, 3.85, SLIDE_W - 2 * MARGIN, 1.5, 20, "DDDDDD")


def slide_title(slide, ctx, spec):
    cover(slide, ctx, spec, 40)


def slide_section(slide, ctx, spec):
    cover(slide, ctx, spec, 34)


def slide_closing(slide, ctx, spec):
    cover(slide, ctx, {"title": spec.get("title", "감사합니다"), "subtitle": spec.get("subtitle")}, 40)


def slide_bullets(slide, ctx, spec):
    header(slide, ctx, spec.get("title", ""))
    items = spec.get("bullets", [])
    box = slide.shapes.add_textbox(Inches(MARGIN), Inches(CONTENT_TOP), Inches(SLIDE_W - 2 * MARGIN),
                                   Inches(CONTENT_BOTTOM - CONTENT_TOP))
    fill_bullets(box, ctx, items, bullet_size(items))


def slide_two_column(slide, ctx, spec):
    header(slide, ctx, spec.get("title", ""))
    col_w = (SLIDE_W - 2 * MARGIN - 0.5) / 2
    for i, key in enumerate(("left", "right")):
        left = MARGIN + i * (col_w + 0.5)
        top = CONTENT_TOP
        heading = spec.get(key + "_title")
        if heading:
            add_text(slide, ctx, heading, left, top, col_w, 0.5, 20, ctx["theme"]["accent"], bold=True)
            top += 0.6
        items = spec.get(key, [])
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(col_w), Inches(CONTENT_BOTTOM - top))
        fill_bullets(box, ctx, items, min(20, bullet_size(items)))


def slide_table(slide, ctx, spec):
    header(slide, ctx, spec.get("title", ""))
    columns, rows = spec["columns"], spec.get("rows", [])
    size = 14 if len(rows) <= 8 else 12 if len(rows) <= 12 else 10
    row_h = 0.3 + size / 72 * 1.2
    height = min(CONTENT_BOTTOM - CONTENT_TOP, row_h * (len(rows) + 1))
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(MARGIN), Inches(CONTENT_TOP),
                                   Inches(SLIDE_W - 2 * MARGIN), Inches(height))
    table = shape.table
    widths = spec.get("col_widths")
    if widths:
        total = sum(widths)
        for i, w in enumerate(widths):
            table.columns[i].width = Inches((SLIDE_W - 2 * MARGIN) * w / total)
    t = ctx["theme"]
    for r, values in enumerate([columns] + rows):
        for c in range(len(columns)):
            cell = table.cell(r, c)
            cell.text = str(values[c]) if c < len(values) else ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(t["primary"] if r == 0 else (t["light"] if r % 2 == 0 else "FFFFFF"))
            for p in cell.text_frame.paragraphs:
                if c > 0 and r > 0 and is_number(cell.text):
                    p.alignment = PP_ALIGN.RIGHT
                for run in p.runs:
                    style_font(run.font, ctx, size, "FFFFFF" if r == 0 else t["text"], bold=(r == 0))


def is_number(text):
    try:
        float(text.replace(",", "").replace("%", ""))
        return True
    except ValueError:
        return False


def slide_chart(slide, ctx, spec):
    header(slide, ctx, spec.get("title", ""))
    chart_type = XL_CHART_TYPE[spec.get("chart_type", "COLUMN_CLUSTERED")]
    if chart_type.name.startswith("XY_"):
        data = XyChartData()
        for s in spec["series"]:
            series = data.add_series(s["name"])
            for x, y in s["points"]:
                series.add_data_point(x, y)
    else:
        data = CategoryChartData(number_format=spec.get("number_format", "General"))
        data.categories = spec["categories"]
        for s in spec["series"]:
            data.add_series(s["name"], s["values"])
    frame = slide.shapes.add_chart(chart_type, Inches(MARGIN), Inches(CONTENT_TOP), Inches(SLIDE_W - 2 * MARGIN),
                                   Inches(CONTENT_BOTTOM - CONTENT_TOP), data)
    chart = frame.chart
    style_font(chart.font, ctx, 14)
    is_pie = "PIE" in chart_type.name or "DOUGHNUT" in chart_type.name
    chart.has_legend = len(spec["series"]) > 1 or is_pie
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if spec.get("data_labels"):
        plot = chart.plots[0]
        plot.has_data_labels = True
        if is_pie:
            plot.data_labels.show_percentage = True
            plot.data_labels.show_value = False
            plot.data_labels.number_format = "0%"
            plot.data_labels.number_format_is_linked = False
    if not is_pie:
        palette = [ctx["theme"]["accent"], ctx["theme"]["primary"], "A5A5A5", "F2A541", "70AD47"]
        for i, series in enumerate(chart.series):
            fmt = series.format
            if "LINE" in chart_type.name or chart_type.name.startswith("XY_"):
                fmt.line.color.rgb = rgb(palette[i % len(palette)])
                fmt.line.width = Pt(2.5)
            else:
                fmt.fill.solid()
                fmt.fill.fore_color.rgb = rgb(palette[i % len(palette)])


def slide_image(slide, ctx, spec):
    header(slide, ctx, spec.get("title", ""))
    path = Path(spec["path"])
    if not path.is_absolute():
        path = ctx["base_dir"] / path
    caption = spec.get("caption")
    box_w = SLIDE_W - 2 * MARGIN
    box_h = CONTENT_BOTTOM - CONTENT_TOP - (0.5 if caption else 0)
    with Image.open(path) as im:
        w_px, h_px = im.size
    scale = min(box_w / w_px, box_h / h_px)
    w, h = w_px * scale, h_px * scale
    left = MARGIN + (box_w - w) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(CONTENT_TOP), Inches(w), Inches(h))
    if caption:
        add_text(slide, ctx, caption, MARGIN, CONTENT_TOP + h + 0.1, box_w, 0.4, 12,
                 ctx["theme"]["muted"], align=PP_ALIGN.CENTER)


BUILDERS = {
    "title": slide_title,
    "section": slide_section,
    "bullets": slide_bullets,
    "two_column": slide_two_column,
    "table": slide_table,
    "chart": slide_chart,
    "image": slide_image,
    "closing": slide_closing,
}
NO_FOOTER = {"title", "section", "closing"}


# ---------------------------------------------------------------- main

def build(spec, out_path, base_dir):
    theme = dict(THEMES[spec.get("theme", "blue")]) if isinstance(spec.get("theme", "blue"), str) \
        else {**THEMES["blue"], **spec["theme"]}
    ctx = {"theme": theme, "font": spec.get("font", DEFAULT_FONT), "footer": spec.get("footer"),
           "base_dir": base_dir}

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    props = prs.core_properties
    props.title = spec.get("title", "")
    props.author = spec.get("author", "")

    for number, s in enumerate(spec["slides"], start=1):
        kind = s.get("type", "bullets")
        if kind not in BUILDERS:
            raise ValueError(f"slide {number}: unknown type '{kind}' (use one of {sorted(BUILDERS)})")
        slide = prs.slides.add_slide(blank)
        BUILDERS[kind](slide, ctx, s)
        if kind not in NO_FOOTER:
            footer(slide, ctx, number)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return len(prs.slides)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("spec", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")  # Korean paths; the Windows console default is cp949

    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    out = args.output or spec.get("output")
    if not out:
        sys.exit("error: no output path; set \"output\" in the spec or pass -o")
    out = Path(out)
    if out.exists():
        sys.exit(f"error: {out} already exists; choose a new file name (e.g. add a date suffix)")

    count = build(spec, out, args.spec.resolve().parent)
    print(f"saved {out} ({count} slides)")


if __name__ == "__main__":
    main()
