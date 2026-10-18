"""Inspect and restructure .pptx files beyond python-pptx's public API.

Usage:
    python pptx_tools.py inspect deck.pptx
    python pptx_tools.py text deck.pptx
    python pptx_tools.py delete deck.pptx 3 -o new.pptx
    python pptx_tools.py move deck.pptx 5 2 -o new.pptx
    python pptx_tools.py duplicate deck.pptx 2 -o new.pptx

Slide numbers are 1-based. Commands that change a deck require -o and refuse
to overwrite an existing file.
"""

import argparse
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


def _check(prs, number):
    if not 1 <= number <= len(prs.slides):
        sys.exit(f"error: slide {number} out of range (deck has {len(prs.slides)} slides)")


def inspect(prs):
    print(f"slides: {len(prs.slides)}  size: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} in")
    for i, slide in enumerate(prs.slides, start=1):
        print(f"\n[{i}] layout={slide.slide_layout.name!r}")
        for shape in slide.shapes:
            kind = shape.shape_type
            detail = ""
            if shape.has_text_frame and shape.text_frame.text.strip():
                detail = repr(shape.text_frame.text.strip().split("\n")[0][:50])
            if getattr(shape, "has_table", False) and shape.has_table:
                t = shape.table
                detail = f"table {len(t.rows)}x{len(t.columns)}"
            if getattr(shape, "has_chart", False) and shape.has_chart:
                c = shape.chart
                detail = f"chart {c.chart_type.name}, {len(list(c.series))} series"
            print(f"  - {kind}: {shape.name} {detail}")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            print(f"  notes: {slide.notes_slide.notes_text_frame.text.strip()[:60]!r}")


def text(prs):
    for i, slide in enumerate(prs.slides, start=1):
        print(f"--- slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if p.text.strip():
                        print("  " * p.level + p.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    print(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            print(f"[notes] {slide.notes_slide.notes_text_frame.text.strip()}")


def delete(prs, number):
    _check(prs, number)
    id_list = prs.slides._sldIdLst
    sld_id = id_list[number - 1]
    prs.part.drop_rel(sld_id.rId)
    id_list.remove(sld_id)


def move(prs, number, to):
    _check(prs, number)
    _check(prs, to)
    id_list = prs.slides._sldIdLst
    sld_id = id_list[number - 1]
    id_list.remove(sld_id)
    id_list.insert(to - 1, sld_id)


def duplicate(prs, number):
    """Append a copy of a slide right after the original.

    Shapes, pictures and hyperlinks are copied; the copy shares image parts with
    the original. Charts and embedded objects are not supported, because two
    slides cannot safely share one chart part.
    """
    _check(prs, number)
    src = prs.slides[number - 1]
    for shape in src.shapes:
        if getattr(shape, "has_chart", False) and shape.has_chart:
            sys.exit("error: duplicating slides with charts is not supported; rebuild the chart instead")
    new = prs.slides.add_slide(src.slide_layout)
    for shape in list(new.shapes):
        shape.element.getparent().remove(shape.element)

    # Re-create the source slide's relationships on the new slide and remember
    # how the rIds map, so references inside the copied XML stay valid.
    rid_map = {}
    for rel in src.part.rels.values():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rid_map[rel.rId] = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            rid_map[rel.rId] = new.part.rels.get_or_add(rel.reltype, rel.target_part)

    r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith(("}nvGrpSpPr", "}grpSpPr")):
            continue
        el_copy = copy.deepcopy(el)
        for node in el_copy.iter():
            for attr, value in node.attrib.items():
                if attr.startswith(r_ns) and value in rid_map:
                    node.set(attr, rid_map[value])
        new.shapes._spTree.append(el_copy)

    if src.has_notes_slide:
        new.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    move(prs, len(prs.slides), number + 1)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=["inspect", "text", "delete", "move", "duplicate"])
    parser.add_argument("deck", type=Path)
    parser.add_argument("numbers", type=int, nargs="*")
    parser.add_argument("-o", "--output", type=Path)
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")  # Korean text; the Windows console default is cp949

    prs = Presentation(args.deck)
    if args.command == "inspect":
        return inspect(prs)
    if args.command == "text":
        return text(prs)

    if not args.output:
        sys.exit("error: -o OUTPUT is required for commands that change the deck")
    if args.output.exists():
        sys.exit(f"error: {args.output} already exists; choose a new file name")
    needed = 2 if args.command == "move" else 1
    if len(args.numbers) != needed:
        sys.exit(f"error: {args.command} takes {needed} slide number(s)")

    {"delete": delete, "move": move, "duplicate": duplicate}[args.command](prs, *args.numbers)
    prs.save(args.output)
    print(f"saved {args.output} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
